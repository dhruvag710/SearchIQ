import sys
from pathlib import Path

import docx

# Add backend directory to sys.path if not present
backend_dir = Path(__file__).resolve().parent.parent.parent / "backend"
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))

from app.ingestion.docx_parser import DocxParser, parse_docx
from app.ingestion.models import Document, Page


def test_docx_parser(tmp_path: Path):
    sample_file = tmp_path / "sample_document.docx"
    doc = docx.Document()
    doc.add_paragraph("First paragraph with UTF-8: 🔥 SearchIQ enterprise test.")
    doc.add_paragraph("")  # empty paragraph
    doc.add_paragraph("   ")  # whitespace-only paragraph
    doc.add_paragraph("Second   paragraph with   extra   spaces.")
    doc.save(sample_file)

    parser = DocxParser()
    parsed_doc = parser.parse(sample_file)

    assert isinstance(parsed_doc, Document)
    assert parsed_doc.document_title == "sample_document"
    assert parsed_doc.total_pages == 1
    assert len(parsed_doc.pages) == 1
    assert parsed_doc.pages[0].page_number == 1

    expected_text = (
        "First paragraph with UTF-8: 🔥 SearchIQ enterprise test.\n\n"
        "Second paragraph with extra spaces."
    )
    assert parsed_doc.pages[0].text == expected_text


def test_parse_docx_helper(tmp_path: Path):
    sample_file = tmp_path / "another_doc.docx"
    doc = docx.Document()
    doc.add_paragraph("Testing parse_docx helper function.")
    doc.save(sample_file)

    parsed_doc = parse_docx(sample_file)

    assert parsed_doc.document_title == "another_doc"
    assert parsed_doc.total_pages == 1
    assert parsed_doc.pages[0].text == "Testing parse_docx helper function."


if __name__ == "__main__":
    import tempfile
    import unittest

    class TestDocxParserUnittest(unittest.TestCase):
        def test_docx_parser(self):
            with tempfile.TemporaryDirectory() as tmpdir:
                path = Path(tmpdir) / "sample_document.docx"
                doc = docx.Document()
                doc.add_paragraph("First paragraph with UTF-8: 🔥 SearchIQ enterprise test.")
                doc.add_paragraph("")
                doc.add_paragraph("   ")
                doc.add_paragraph("Second   paragraph with   extra   spaces.")
                doc.save(path)

                parser = DocxParser()
                parsed_doc = parser.parse(path)

                self.assertEqual(parsed_doc.document_title, "sample_document")
                self.assertEqual(parsed_doc.total_pages, 1)
                self.assertEqual(parsed_doc.pages[0].page_number, 1)
                expected_text = (
                    "First paragraph with UTF-8: 🔥 SearchIQ enterprise test.\n\n"
                    "Second paragraph with extra spaces."
                )
                self.assertEqual(parsed_doc.pages[0].text, expected_text)

        def test_parse_docx_helper(self):
            with tempfile.TemporaryDirectory() as tmpdir:
                path = Path(tmpdir) / "another_doc.docx"
                doc = docx.Document()
                doc.add_paragraph("Testing parse_docx helper function.")
                doc.save(path)

                parsed_doc = parse_docx(path)
                self.assertEqual(parsed_doc.document_title, "another_doc")
                self.assertEqual(parsed_doc.total_pages, 1)
                self.assertEqual(parsed_doc.pages[0].text, "Testing parse_docx helper function.")

    unittest.main()
