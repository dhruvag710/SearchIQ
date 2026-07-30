from pathlib import Path

import docx

from app.ingestion.base import DocumentParser
from app.ingestion.cleaner import clean_text
from app.ingestion.models import Document, Page


class DocxParser(DocumentParser):
    """Extract structured text from Microsoft Word (.docx) documents."""

    def parse(self, path: Path) -> Document:
        """Read a DOCX document and return structured text for the document."""
        doc = docx.Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        joined_text = "\n\n".join(paragraphs)

        return Document(
            document_title=path.stem,
            total_pages=1,
            pages=[
                Page(
                    page_number=1,
                    text=clean_text(joined_text),
                )
            ],
        )


def parse_docx(path: Path) -> Document:
    """Read a DOCX document and return structured text for the document."""
    return DocxParser().parse(path)
